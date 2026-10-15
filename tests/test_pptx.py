"""Markdown → PPTX, в том числе поверх шаблона организации."""
from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from latex_docx import PptxExportError, markdown_to_pptx
from latex_docx.cli import main
from latex_docx.pptx_export import parse_slide_markdown

SLIDES = """# Доклад
Лаборатория

---

## Результаты

- первый
- второй

???
Сказать про Q.

---

## Схема

Текст слайда.

![схема](missing.png)
"""


def test_parse_title_content_and_notes(tmp_path: Path):
    slides = parse_slide_markdown(SLIDES, tmp_path)
    assert len(slides) == 3
    assert slides[0].kind == "title"
    assert slides[0].title == "Доклад"
    assert slides[0].subtitle == "Лаборатория"
    assert slides[1].lines == ["первый", "второй"]
    assert slides[1].notes == "Сказать про Q."
    assert slides[2].images[0].name == "missing.png"


def test_builds_pptx_and_keeps_missing_image_text(tmp_path: Path):
    md = tmp_path / "talk.md"
    md.write_text(SLIDES, encoding="utf-8")
    out = markdown_to_pptx(md, tmp_path / "talk.pptx")
    assert out.name == "out_talk.pptx"
    prs = Presentation(str(out))
    assert len(prs.slides) == 3
    assert prs.slides[0].shapes.title.text == "Доклад"
    assert "первый" in prs.slides[1].placeholders[1].text
    assert prs.slides[1].notes_slide.notes_text_frame.text == "Сказать про Q."
    assert "missing.png" in prs.slides[2].placeholders[1].text


def test_template_sets_size_and_drops_sample_slides(tmp_path: Path):
    template = tmp_path / "org.pptx"
    base = Presentation()
    base.slide_width = Inches(10)
    base.slide_height = Inches(5.625)
    base.slides.add_slide(base.slide_layouts[0])
    base.save(template)

    md = tmp_path / "talk.md"
    md.write_text("# Только титул\n\n## Дальше\n\nАбзац.\n", encoding="utf-8")
    out = markdown_to_pptx(
        md,
        tmp_path / "deck.pptx",
        template=template,
        content_layout="Title Only",
    )
    prs = Presentation(str(out))
    assert prs.slide_width == Inches(10)
    assert len(prs.slides) == 2
    assert prs.slides[0].slide_layout.name == "Title Slide"
    assert prs.slides[1].slide_layout.name == "Title Only"
    assert prs.slides[1].shapes.title.text == "Дальше"


def test_unknown_layout_names_the_ones_present(tmp_path: Path):
    md = tmp_path / "talk.md"
    md.write_text("# Заголовок\n", encoding="utf-8")
    template = tmp_path / "org.pptx"
    Presentation().save(template)
    with pytest.raises(PptxExportError) as exc:
        markdown_to_pptx(md, tmp_path / "x.pptx", template=template, title_layout="Нет такого")
    assert "Title Slide" in str(exc.value)


def test_cli_pptx(tmp_path: Path):
    md = tmp_path / "talk.md"
    md.write_text("# Привет\n", encoding="utf-8")
    main([str(md), "-o", str(tmp_path / "deck.pptx"), "--no-log"])
    assert (tmp_path / "out_deck.pptx").exists()

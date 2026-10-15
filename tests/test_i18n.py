"""Language of labels inserted into the document."""
from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from latex_docx.docx_builder import DocxBuilder
from latex_docx.i18n import normalize_lang
from latex_docx.models import Figure, Paper, Run, TableBlock


def test_normalize_lang_aliases_and_reject():
    assert normalize_lang("EN") == "en"
    assert normalize_lang("russian") == "ru"
    assert normalize_lang(None) == "ru"
    with pytest.raises(ValueError):
        normalize_lang("de")


def _texts(builder: DocxBuilder) -> str:
    return "\n".join(p.text for p in builder.doc.paragraphs)


def test_default_captions_stay_russian():
    paper = Paper(
        udc="004",
        title_ru="Заголовок",
        references=["1. Источник"],
        authors_info=["Автор"],
        body=[
            Figure(paths=[Path("missing.png")], caption="схема", number=1),
            TableBlock(rows=[[[Run("ячейка")]]], caption="данные", number=2),
        ],
    )
    builder = DocxBuilder()
    builder.build(paper, template="plain")
    text = _texts(builder)
    assert "УДК: 004" in text
    assert "Рис. 1. схема" in text
    assert "Таблица 2. данные" in text
    assert "Список литературы" in text
    assert "Сведения об авторах" in text
    assert "[рисунок не найден: missing.png]" in text


def test_english_plain_uses_english_labels_and_falls_back_to_title_ru():
    paper = Paper(
        udc="004",
        title_ru="Note",
        abstract_ru="Short note.",
        keywords_ru="one, two",
        references=["1. Source"],
        body=[Figure(paths=[Path("missing.png")], caption="layout", number=1)],
    )
    builder = DocxBuilder(lang="en")
    builder.build(paper, template="plain")
    text = _texts(builder)
    assert "Note" in text
    assert "Abstract. Short note." in text
    assert "Keywords: one, two" in text
    assert "UDC: 004" in text
    assert "Fig. 1. layout" in text
    assert "References" in text
    assert "[image not found: missing.png]" in text


def test_tdj_keeps_both_front_matter_blocks():
    paper = Paper(title_ru="Русский", title_en="English", abstract_ru="А.", abstract_en="A.")
    builder = DocxBuilder(lang="en")
    builder.build(paper, template="tdj")
    text = _texts(builder)
    assert "Русский" in text and "English" in text
    assert "Аннотация. А." in text
    assert "Abstract. A." in text


def test_caption_already_numbered_is_not_doubled():
    paper = Paper(body=[Figure(paths=[Path("a.png")], caption="Fig. 3. already", number=3)])
    builder = DocxBuilder(lang="ru")
    builder.build(paper, template="plain")
    assert "Fig. 3. already" in _texts(builder)
    assert "Рис. 3. Fig." not in _texts(builder)


def test_english_plain_prefers_english_title():
    paper = Paper(title_ru="Русский", title_en="English")
    builder = DocxBuilder(lang="en")
    builder.build(paper, template="plain")
    text = _texts(builder)
    assert "English" in text
    assert "Русский" not in text


def test_roundtrip_saved_docx(tmp_path):
    paper = Paper(body=[TableBlock(rows=[[[Run("cell")]]], caption="rows", number=1)])
    builder = DocxBuilder(lang="en")
    builder.build(paper, template="plain")
    path = tmp_path / "out_lang.docx"
    builder.doc.save(str(path))
    saved = "\n".join(p.text for p in Document(str(path)).paragraphs)
    assert "Table 1. rows" in saved

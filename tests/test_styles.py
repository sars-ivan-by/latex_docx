"""Heading styles, organization DOCX templates, lists, and footnotes."""
from __future__ import annotations

import io
import zipfile
from pathlib import Path

import docx
import pytest
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from latex_docx.docx_builder import DocxBuilder
from latex_docx.md_parser import MarkdownParser
from latex_docx.models import Heading, Paper, Para, Run


def _style_id(paragraph) -> str:
    return paragraph.style.element.get(qn("w:styleId")) or ""


def test_headings_use_outline_styles():
    paper = Paper(body=[Heading("Введение", level=1), Heading("Подраздел", level=2)])
    builder = DocxBuilder(template="article")
    builder.build(paper, template="article")
    headings = [p for p in builder.doc.paragraphs if p.text in {"Введение", "Подраздел"}]
    assert [_style_id(p) for p in headings] == ["Heading1", "Heading2"]
    assert headings[0].runs[0].font.name == "Times New Roman"
    assert headings[0].runs[0].font.size == Pt(14)
    outline = builder.doc.styles["Heading 1"].element.find(qn("w:pPr")).find(qn("w:outlineLvl"))
    assert outline.get(qn("w:val")) == "0"


def test_style_template_keeps_page_and_drops_sample(tmp_path: Path):
    org = tmp_path / "org.docx"
    source = docx.Document()
    source.styles["Normal"].font.name = "Courier New"
    source.sections[0].page_width = Inches(6)
    source.add_paragraph("SAMPLE SHOULD GO")
    source.save(org)

    paper = Paper(
        title_ru="Title",
        body=[Heading("Section", level=1), Para(runs=[Run("body text")])],
    )
    builder = DocxBuilder(style_template=org)
    builder.build(paper, template="plain")
    text = "\n".join(p.text for p in builder.doc.paragraphs)
    assert "SAMPLE SHOULD GO" not in text
    assert builder.doc.sections[0].page_width == Inches(6)
    assert builder.doc.styles["Normal"].font.name == "Courier New"
    body = next(p for p in builder.doc.paragraphs if p.text == "body text")
    assert body.runs[0].font.name is None
    heading = next(p for p in builder.doc.paragraphs if p.text == "Section")
    assert _style_id(heading) == "Heading1"


def test_dotx_template_opens(tmp_path: Path):
    docx_path = tmp_path / "org.docx"
    dotx_path = tmp_path / "org.dotx"
    source = docx.Document()
    source.sections[0].page_width = Inches(7)
    source.add_paragraph("gone")
    source.save(docx_path)
    buf = io.BytesIO()
    with zipfile.ZipFile(docx_path) as zin, zipfile.ZipFile(buf, "w") as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "[Content_Types].xml":
                data = data.replace(
                    b"wordprocessingml.document.main+xml",
                    b"wordprocessingml.template.main+xml",
                )
            zout.writestr(info, data)
    dotx_path.write_bytes(buf.getvalue())

    builder = DocxBuilder(style_template=dotx_path)
    builder.build(Paper(body=[Para(runs=[Run("kept")])]), template="article")
    assert builder.doc.sections[0].page_width == Inches(7)
    assert "kept" in "\n".join(p.text for p in builder.doc.paragraphs)


def test_article_template_is_one_block():
    paper = Paper(title_ru="Русский", title_en="English")
    builder = DocxBuilder(lang="en")
    builder.build(paper, template="article")
    text = "\n".join(p.text for p in builder.doc.paragraphs)
    assert "English" in text
    assert "Русский" not in text


def test_unknown_template_suffix(tmp_path: Path):
    bad = tmp_path / "org.txt"
    bad.write_text("nope", encoding="utf-8")
    with pytest.raises(ValueError):
        DocxBuilder(style_template=bad)


def test_nested_lists_and_footnotes(tmp_path: Path):
    md = tmp_path / "note.md"
    md.write_text(
        """# Note

See the source[^src].

- outer
  - inner
1. first

[^src]: The laboratory log.
""",
        encoding="utf-8",
    )
    paper = MarkdownParser(md).parse()
    lists = [b for b in paper.body if b.__class__.__name__ == "ListBlock"]
    assert [(b.level, b.ordered, b.items[0][0].text) for b in lists] == [
        (0, False, "outer"),
        (1, False, "inner"),
        (0, True, "first"),
    ]
    note = next(r for b in paper.body if isinstance(b, Para) for r in b.runs if r.footnote)
    assert note.footnote == "The laboratory log."

    builder = DocxBuilder(lang="en")
    builder.build(paper, template="article")
    path = tmp_path / "out_note.docx"
    builder.doc.save(path)
    with zipfile.ZipFile(path) as zf:
        footnotes = zf.read("word/footnotes.xml").decode("utf-8")
    assert "The laboratory log." in footnotes
    assert "footnoteReference" in builder.doc.paragraphs[1]._p.xml

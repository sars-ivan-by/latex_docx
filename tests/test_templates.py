"""Document templates differ in type, alignment, and front matter."""
from __future__ import annotations

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from latex_docx.docx_builder import DocxBuilder
from latex_docx.models import Paper, Para, Run


def _built(template: str) -> DocxBuilder:
    paper = Paper(
        udc="004",
        title_ru="Hello",
        title_en="Hi",
        abstract_ru="Short.",
        body=[Para(runs=[Run("Body text.")])],
    )
    builder = DocxBuilder(template=template)
    builder.build(paper, template=template)
    return builder


def _body(builder: DocxBuilder):
    return next(p for p in builder.doc.paragraphs if p.text == "Body text.")


def test_note_is_a_plain_office_page():
    builder = _built("note")
    assert builder.doc.sections[0].left_margin == Inches(1)
    assert _body(builder).runs[0].font.name == "Calibri"
    assert _body(builder).runs[0].font.size == Pt(11)
    assert _body(builder).alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert "УДК" in "\n".join(p.text for p in builder.doc.paragraphs)
    assert "Hi" not in "\n".join(p.text for p in builder.doc.paragraphs)


def test_letter_drops_journal_labels_and_keeps_the_prose():
    builder = _built("letter")
    text = "\n".join(p.text for p in builder.doc.paragraphs)
    assert "УДК" not in text
    assert "Аннотация" not in text
    assert "Short." in text
    assert "Hi" not in text
    title = next(p for p in builder.doc.paragraphs if p.text == "Hello")
    assert title.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert builder.doc.sections[0].left_margin == Inches(1.15)


def test_memo_is_tighter_than_a_letter():
    letter = _built("letter")
    memo = _built("memo")
    assert memo.doc.sections[0].left_margin == Inches(0.85)
    assert memo.doc.sections[0].left_margin < letter.doc.sections[0].left_margin
    assert "Аннотация" not in "\n".join(p.text for p in memo.doc.paragraphs)


def test_report_is_justified_times():
    builder = _built("report")
    body = _body(builder)
    assert body.runs[0].font.name == "Times New Roman"
    assert body.runs[0].font.size == Pt(12)
    assert body.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert body.paragraph_format.first_line_indent == Inches(0.3)


def test_article_stays_fourteen_point_and_tdj_stays_bilingual():
    article = _built("article")
    assert _body(article).runs[0].font.size == Pt(14)
    assert "Hi" not in "\n".join(p.text for p in article.doc.paragraphs)
    tdj = _built("tdj")
    text = "\n".join(p.text for p in tdj.doc.paragraphs)
    assert "Hello" in text and "Hi" in text
    assert "Аннотация" in text

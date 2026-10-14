"""Light Markdown → DOCX."""
from __future__ import annotations

from pathlib import Path

from docx import Document

from latex_docx import convert
from latex_docx.cli import main
from latex_docx.convert import markdown_to_docx
from latex_docx.md_parser import CODE_FONT, MarkdownParser, parse_md_inline
from latex_docx.models import Equation, Figure, Heading, ListBlock, Para, TableBlock


def test_inline_bold_italic_code_and_link():
    runs = parse_md_inline("see **bold** and *it* and `x` plus [site](https://example.com)")
    assert runs[0].text == "see "
    assert runs[1].text == "bold" and runs[1].bold
    assert any(r.text == "it" and r.italic for r in runs)
    code = next(r for r in runs if r.font == CODE_FONT)
    assert code.text == "x"
    assert any("example.com" in (r.text or "") for r in runs)


def test_inline_underscore_inside_word_stays():
    runs = parse_md_inline("var_name_here and _em_")
    assert runs[0].text == "var_name_here and "
    assert runs[1].text == "em" and runs[1].italic


def test_parser_title_lists_table_code_math(tmp_path: Path):
    md = tmp_path / "note.md"
    md.write_text(
        """# Записка

Абзац с $E=mc^2$.

## Раздел

- первый
- [x] сделано

1. раз
2. два

| a | b |
| --- | --- |
| 1 | 2 |

```
print(1)
```

> цитата

$$
x=1
$$

![схема](missing.png)
""",
        encoding="utf-8",
    )
    paper = MarkdownParser(md).parse()
    assert paper.title_ru == "Записка"
    assert any(isinstance(b, Heading) and b.text == "Раздел" for b in paper.body)
    lists = [b for b in paper.body if isinstance(b, ListBlock)]
    assert lists[0].ordered is False
    assert lists[0].items[1][0].text.startswith("☑")
    assert lists[1].ordered is True
    tables = [b for b in paper.body if isinstance(b, TableBlock)]
    assert tables[0].rows[1][0][0].text == "1"
    code = next(b for b in paper.body if isinstance(b, Para) and b.runs and b.runs[0].font == CODE_FONT)
    assert "print(1)" in code.runs[0].text
    assert any(isinstance(b, Equation) and "x=1" in b.latex for b in paper.body)
    fig = next(b for b in paper.body if isinstance(b, Figure))
    assert fig.caption == "схема"
    assert fig.paths[0].name == "missing.png"
    quote = next(b for b in paper.body if isinstance(b, Para) and b.left_indent)
    assert quote.runs[0].italic
    assert "цитата" in quote.runs[0].text


def test_yaml_title_keeps_h1(tmp_path: Path):
    md = tmp_path / "t.md"
    md.write_text("---\ntitle: Из YAML\n---\n\n# Раздел\n\nТекст.\n", encoding="utf-8")
    paper = MarkdownParser(md).parse()
    assert paper.title_ru == "Из YAML"
    assert any(isinstance(b, Heading) and b.text == "Раздел" for b in paper.body)


def test_markdown_to_docx_and_cli(tmp_path: Path):
    md = tmp_path / "letter.md"
    md.write_text("# Письмо\n\nТекст **важен**.\n\n`код`\n", encoding="utf-8")
    out = markdown_to_docx(md, tmp_path / "letter.docx")
    assert out.name == "out_letter.docx"
    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs]
    assert "Письмо" in texts
    assert any("важен" in t for t in texts)
    fonts = [r.font.name for p in doc.paragraphs for r in p.runs]
    assert CODE_FONT in fonts

    cli_out = tmp_path / "via.docx"
    main([str(md), "-o", str(cli_out), "--no-log"])
    assert (tmp_path / "out_via.docx").exists()


def test_convert_dispatches_markdown_and_rejects_unknown(tmp_path: Path):
    md = tmp_path / "n.md"
    md.write_text("# Заголовок\n\nТекст.\n", encoding="utf-8")
    out = convert(md, tmp_path / "n.docx")
    assert out.name == "out_n.docx"
    assert out.exists()
    bad = tmp_path / "n.txt"
    bad.write_text("x", encoding="utf-8")
    try:
        convert(bad)
    except ValueError as exc:
        assert ".txt" in str(exc)
    else:
        raise AssertionError("expected ValueError")

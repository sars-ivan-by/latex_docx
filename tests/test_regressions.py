"""Regression tests for latex_docx (cases, \\_, eqref, align, sec refs)."""
from __future__ import annotations

from pathlib import Path

from docx.oxml.ns import qn

from latex_docx.authors import apply_authors_db, load_authors_db
from latex_docx.constants import FONT
from latex_docx.docx_builder import DocxBuilder
from latex_docx.inline import parse_inline
from latex_docx.math_omml import _prep_latex_math, latex_math_to_text
from latex_docx.models import Paper, Para
from latex_docx.parser import LatexParser

FIXTURES = Path(__file__).resolve().parent / "fixtures"
AUTHORS_DB = FIXTURES / "authors.json"


def test_prep_keeps_cases_row_separators():
    out = _prep_latex_math(r"\begin{cases} a & b \\ c & d \end{cases}")
    assert r"\begin{cases}" in out
    assert r"\\" in out
    assert "&" in out
    assert "\0R\0" not in out


def test_prep_align_becomes_aligned_not_flattened():
    out = _prep_latex_math(r"\begin{align} a &= b \\ c &= d \end{align}")
    assert r"\begin{aligned}" in out
    assert r"\\" in out
    assert "a = b c = d" not in out.replace(" ", "")


def test_mathrm_underscore_not_subscript():
    assert latex_math_to_text(r"\mathrm{exp\_fit}") == "exp_fit"
    assert r"\_" in _prep_latex_math(r"\mathrm{exp\_fit}") or "exp" in _prep_latex_math(
        r"\mathrm{exp\_fit}"
    )


def test_eqref_stays_separate_run():
    runs = parse_inline(r"see \eqref{eq:foo} and more")
    ref = next(r for r in runs if r.ref_key)
    assert ref.ref_key == "eq:foo"
    assert ref.eqref is True
    assert ref.text == "eq:foo"
    assert runs[0].text == "see "
    assert runs[-1].text == " and more"


def test_sec_and_eqref_resolve(tmp_path: Path):
    tex = tmp_path / "sample.tex"
    tex.write_text(
        r"""
\documentclass{article}
\begin{document}
\section{Quality}\label{sec:quality}
\subsection{Details}\label{sec:details}
See \ref{sec:quality} and \ref{sec:details}.
\begin{equation}\label{eq:one}
x=1
\end{equation}
See \eqref{eq:one}.
\end{document}
""",
        encoding="utf-8",
    )
    paper = LatexParser(tex).parse()
    texts = [
        "".join(r.text for r in b.runs) for b in paper.body if isinstance(b, Para)
    ]
    joined = " ".join(texts)
    assert "See 1 and 1.1." in joined
    assert "See (1)." in joined


def test_apply_authors_preserves_careful_tex():
    paper = Paper(
        authors_ru="И. А. Сукач¹",
        affiliations_ru="¹ Длинная организация из tex — более двадцати символов",
    )
    apply_authors_db(paper, load_authors_db(AUTHORS_DB))
    assert paper.authors_ru == "И. А. Сукач¹"
    assert paper.affiliations_ru.startswith("¹")
    assert paper.authors_en


def test_apply_authors_adds_superscripts_when_tex_has_long_aff_without_marks():
    """Long affiliation text without ¹… must not block authors.json marks."""
    paper = Paper(
        authors_ru="Сукач Иван Андреевич, Меркулов Андрей Александрович",
        affiliations_ru=(
            "Аспирант, ассистент\n"
            "Национальный исследовательский университет «Московский энергетический институт»\n"
            "Москва, Россия"
        ),
    )
    apply_authors_db(paper, load_authors_db(AUTHORS_DB))
    assert "¹" in paper.authors_ru
    assert "²" in paper.authors_ru
    assert paper.affiliations_ru.startswith("¹")
    assert "²" in paper.affiliations_ru


def test_apply_authors_explicit_overwrites():
    paper = Paper(
        authors_ru="И. А. Сукач¹",
        affiliations_ru="¹ Длинная организация из tex — более двадцати символов",
    )
    apply_authors_db(
        paper, load_authors_db(AUTHORS_DB), author_ids=["sukach", "merkulov"]
    )
    assert "Меркулов" in paper.authors_ru
    assert "Sukach" in paper.authors_en or "I. A." in paper.authors_en


def test_table_caption_before_table_right_aligned():
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from latex_docx.models import Run

    b = DocxBuilder()
    b.add_table([[[Run("a")], [Run("b")]]], caption="Параметры", number=2)
    body = list(b.doc.element.body)
    tags = [child.tag.split("}")[-1] for child in body]
    assert tags.index("p") < tags.index("tbl")
    cap = b.doc.paragraphs[0]
    assert cap.text == "Таблица 2. Параметры"
    assert cap.alignment == WD_ALIGN_PARAGRAPH.RIGHT


def test_run_font_sets_rFonts_ascii_hansi():
    builder = DocxBuilder()
    p = builder.add_p("hello")
    run = p.runs[0]
    r_pr = run._element.find(qn("w:rPr"))
    assert r_pr is not None
    r_fonts = r_pr.find(qn("w:rFonts"))
    assert r_fonts is not None
    face = builder.profile.font
    assert r_fonts.get(qn("w:ascii")) == face
    assert r_fonts.get(qn("w:hAnsi")) == face


def test_textbf_preserves_inner_math():
    runs = parse_inline(r"\textbf{hello $x^2$ world}")
    math_runs = [r for r in runs if r.math]
    assert len(math_runs) == 1
    assert math_runs[0].math == "x^2"
    assert all(r.bold for r in runs)


def test_tabular_cells_keep_math_runs():
    from latex_docx.inline import parse_tabular

    rows = parse_tabular(r"{cc} a & $x^2$ \\ b & c")
    assert rows[0][1][0].math == "x^2"


def test_booktabs_multicolumn():
    from latex_docx.inline import parse_tabular

    rows = parse_tabular(
        r"""{lcc}
\toprule
A & \multicolumn{2}{c}{Head} \\
\midrule
1 & $a$ & b \\
\bottomrule
"""
    )
    assert len(rows) == 2
    assert any(r.math == "a" for r in rows[1][1])


def test_bmod_spacing_in_flat_text():
    t = latex_math_to_text(r"a \bmod b")
    assert "mod" in t
    assert t.index("a") < t.index("mod") < t.index("b")


def test_unwrap_omml_boxes_with_boxpr():
    from latex_docx.math_omml import _unwrap_omml_boxes

    xml = (
        "<m:box><m:boxPr><m:aln/></m:boxPr>"
        "<m:e><m:r><m:t>x</m:t></m:r></m:e></m:box>"
    )
    out = _unwrap_omml_boxes(xml)
    assert "m:box" not in out
    assert "x" in out


def test_vspace_unbalanced_does_not_raise():
    runs = parse_inline(r"before \vspace{1em after")
    assert runs  # no exception


def test_expand_input(tmp_path: Path):
    (tmp_path / "child.tex").write_text("CHILDTEXT", encoding="utf-8")
    main = tmp_path / "main.tex"
    main.write_text(
        r"""
\documentclass{article}
\begin{document}
\input{child}
\end{document}
""",
        encoding="utf-8",
    )
    paper = LatexParser(main).parse()
    texts = [
        "".join(r.text for r in b.runs) for b in paper.body if isinstance(b, Para)
    ]
    assert any("CHILDTEXT" in t for t in texts)


def test_bbl_bibliography(tmp_path: Path):
    tex = tmp_path / "paper.tex"
    tex.write_text(
        r"""
\documentclass{article}
\begin{document}
Cite \cite{abc}.
\bibliography{refs}
\end{document}
""",
        encoding="utf-8",
    )
    (tmp_path / "paper.bbl").write_text(
        r"""
\begin{thebibliography}{1}
\bibitem{abc} Author. Title. 2020.
\end{thebibliography}
""",
        encoding="utf-8",
    )
    paper = LatexParser(tex).parse()
    assert paper.cite_map["abc"] == 1
    assert paper.references[0].startswith("1. Author")
    assert "thebibliography" not in paper.references[0]


def test_cite_in_abstract_uses_cite_map(tmp_path: Path):
    tex = tmp_path / "a.tex"
    tex.write_text(
        r"""
\documentclass{article}
\begin{document}
\begin{abstract}
See \cite{foo}.
\end{abstract}
\begin{thebibliography}{1}
\bibitem{foo} Foo.
\end{thebibliography}
\end{document}
""",
        encoding="utf-8",
    )
    paper = LatexParser(tex).parse()
    runs = parse_inline(paper.abstract_ru_tex, paper.cite_map)
    assert any("[1]" in (r.text or "") for r in runs)


def test_tall_numbered_equation_uses_table():
    from latex_docx.models import Equation

    b = DocxBuilder()
    b.add_equation(
        Equation(
            latex=r"\begin{cases} a & b \\ c & d \end{cases}",
            number=1,
            text="cases",
        )
    )
    assert len(b.doc.tables) == 1


def test_diff_sees_formula_edits(tmp_path: Path):
    from latex_docx.diff import diff_docx_paragraphs
    from latex_docx.math_omml import make_equation

    b1 = DocxBuilder()
    b1.add_equation(make_equation("x=1", number=1))
    p1 = tmp_path / "a.docx"
    b1.doc.save(str(p1))
    b2 = DocxBuilder()
    b2.add_equation(make_equation("x=2", number=1))
    p2 = tmp_path / "b.docx"
    b2.doc.save(str(p2))
    changes = diff_docx_paragraphs(p1, p2)
    assert changes
    assert "x=1" in changes[0]["base"]
    assert "x=2" in changes[0]["edited"]


def test_resolve_by_agent_env(monkeypatch):
    from latex_docx.util import resolve_by

    monkeypatch.delenv("LATEX_DOCX_BY", raising=False)
    monkeypatch.setenv("CURSOR_AGENT", "1")
    assert resolve_by() == "agent"
    assert resolve_by("user") == "user"


def test_multi_includegraphics_one_number(tmp_path: Path):
    tex = tmp_path / "f.tex"
    # Missing image files are fine — we check numbering
    tex.write_text(
        r"""
\documentclass{article}
\begin{document}
\begin{figure}
\includegraphics{a.png}
\includegraphics{b.png}
\caption{Two panels}
\label{fig:two}
\end{figure}
See \ref{fig:two}.
\end{document}
""",
        encoding="utf-8",
    )
    paper = LatexParser(tex).parse()
    from latex_docx.models import Figure

    figs = [b for b in paper.body if isinstance(b, Figure)]
    assert len(figs) == 1
    assert figs[0].number == 1
    assert len(figs[0].paths) == 2

"""DOCX paragraph diff and markdown report."""
from __future__ import annotations

import re
from pathlib import Path
from typing import List, Optional, Union

import docx

from .constants import DEFAULT_DIFF_REPORT, OMML_NS
from .util import ensure_out_prefix


def _omml_texts(element) -> List[str]:
    """Collect visible text from OMML (m:t) under a paragraph/table cell element."""
    texts: List[str] = []
    for node in element.iter(f"{{{OMML_NS}}}t"):
        if node.text:
            texts.append(node.text)
    return texts


def _para_visible_text(p) -> str:
    """Paragraph text including OMML math (p.text alone often yields only '(1)')."""
    plain = (p.text or "").strip()
    omml_parts = _omml_texts(p._p)
    if not omml_parts:
        return plain
    # Deduplicate while preserving order
    seen = set()
    unique = []
    for t in omml_parts:
        if t not in seen:
            seen.add(t)
            unique.append(t)
    math_blob = "".join(unique)
    if not plain:
        return math_blob
    if math_blob and math_blob not in plain:
        # Numbered eq paragraphs: tabs + (N) in plain, formula in OMML
        return f"{math_blob} {plain}".strip()
    return plain


def _docx_para_texts(path: Union[str, Path]) -> List[str]:
    """Non-empty paragraph (+ table cell) texts; include OMML content."""
    doc = docx.Document(str(path))
    out: List[str] = []
    for p in doc.paragraphs:
        t = _para_visible_text(p)
        if t:
            out.append(t)
    # Equations rendered as borderless tables also need coverage
    for table in doc.tables:
        for row in table.rows:
            cell_texts = []
            for cell in row.cells:
                for p in cell.paragraphs:
                    t = _para_visible_text(p)
                    if t:
                        cell_texts.append(t)
            joined = " | ".join(cell_texts).strip()
            if joined:
                out.append(joined)
    return out



def _norm_for_match(text: str) -> str:
    """Normalize for fuzzy matching: whitespace, dashes, NBSP."""
    t = text.replace("\u00a0", " ").replace("–", "-").replace("—", "-").replace("−", "-")
    t = re.sub(r"\s+", " ", t).strip().casefold()
    return t



def diff_docx_paragraphs(base_docx: Union[str, Path], edited_docx: Union[str, Path]) -> List[dict]:
    """Diff DOCX by content alignment (SequenceMatcher), not by paragraph index.

    Index-based diff breaks when one extra/missing paragraph shifts the rest of
    the document — almost everything looks changed. Here paragraphs are matched
    by text similarity so an abstract-only edit stays one change.
    """
    from difflib import SequenceMatcher

    base = _docx_para_texts(base_docx)
    edited = _docx_para_texts(edited_docx)
    base_n = [_norm_for_match(t) for t in base]
    edited_n = [_norm_for_match(t) for t in edited]

    sm = SequenceMatcher(a=base_n, b=edited_n, autojunk=False)
    changes: List[dict] = []

    def emit_replace(bi: int, bj: int, ei: int, ej: int) -> None:
        b_chunk, e_chunk = base[bi:bj], edited[ei:ej]
        b_n, e_n = base_n[bi:bj], edited_n[ei:ej]
        # Align inside the replace hunk so 1:1 edits stay paired
        sub = SequenceMatcher(a=b_n, b=e_n, autojunk=False)
        for stag, si1, si2, sj1, sj2 in sub.get_opcodes():
            if stag == "equal":
                continue
            if stag == "replace" and (si2 - si1) == (sj2 - sj1):
                for k in range(si2 - si1):
                    changes.append(
                        {
                            "kind": "changed",
                            "base_index": bi + si1 + k,
                            "edited_index": ei + sj1 + k,
                            "base": b_chunk[si1 + k],
                            "edited": e_chunk[sj1 + k],
                        }
                    )
            else:
                for k in range(si1, si2):
                    changes.append(
                        {
                            "kind": "removed",
                            "base_index": bi + k,
                            "edited_index": None,
                            "base": b_chunk[k],
                            "edited": "",
                        }
                    )
                for k in range(sj1, sj2):
                    changes.append(
                        {
                            "kind": "added",
                            "base_index": None,
                            "edited_index": ei + k,
                            "base": "",
                            "edited": e_chunk[k],
                        }
                    )

    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            continue
        if tag == "replace":
            emit_replace(i1, i2, j1, j2)
        elif tag == "delete":
            for k in range(i1, i2):
                changes.append(
                    {
                        "kind": "removed",
                        "base_index": k,
                        "edited_index": None,
                        "base": base[k],
                        "edited": "",
                    }
                )
        elif tag == "insert":
            for k in range(j1, j2):
                changes.append(
                    {
                        "kind": "added",
                        "base_index": None,
                        "edited_index": k,
                        "base": "",
                        "edited": edited[k],
                    }
                )
    return changes



def write_docx_diff_report(
    base_docx: Union[str, Path],
    edited_docx: Union[str, Path],
    report_path: Optional[Union[str, Path]] = None,
) -> Path:
    changes = diff_docx_paragraphs(base_docx, edited_docx)
    report_path = ensure_out_prefix(report_path) if report_path else DEFAULT_DIFF_REPORT
    kind_en = {"changed": "changed", "removed": "removed", "added": "added"}
    lines = [
        f"# DOCX diff: `{Path(edited_docx).name}` vs `{Path(base_docx).name}`",
        "",
        f"Changes: **{len(changes)}** "
        "(matched by paragraph text; empty paragraphs are ignored).",
        "",
        "A full DOCX→TeX import is unreliable (OMML formulas, tables, fields).",
        "Copy the edits below back into the `.tex` or `.md` by hand.",
        "",
        "`base` should be a fresh conversion of the same source (`make convert`),",
        "and `edited` should be that file after the Word edits.",
        "Comparing against an older DOCX from another template adds noise.",
        "",
    ]
    for n, ch in enumerate(changes, 1):
        kind = ch.get("kind", "changed")
        bi, ei = ch.get("base_index"), ch.get("edited_index")
        loc = []
        if bi is not None:
            loc.append(f"base¶{bi}")
        if ei is not None:
            loc.append(f"edited¶{ei}")
        lines.append(f"## {n}. {kind_en.get(kind, kind)} ({', '.join(loc) or '—'})")
        lines.append("")
        if kind in ("changed", "removed"):
            lines.append("**before:**")
            lines.append("")
            lines.append(ch["base"] or "_(empty)_")
            lines.append("")
        if kind in ("changed", "added"):
            lines.append("**after:**")
            lines.append("")
            lines.append(ch["edited"] or "_(empty)_")
            lines.append("")
    if not changes:
        lines.append("_No differences._")
        lines.append("")
    report_path.write_text("\n".join(lines), encoding="utf-8")
    return report_path

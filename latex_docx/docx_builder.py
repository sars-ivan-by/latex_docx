"""python-docx TD-J builder."""
from __future__ import annotations

import re
from pathlib import Path
from typing import List, Optional, Union

import docx
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

from .i18n import CAPTION_SKIP, labels_for, normalize_lang
from .templates import profile_for
from .constants import (
    BODY_PT,
    FIG_WIDTH,
    FIRST_INDENT,
    FONT,
    MARGIN,
    OMML_NS,
    PAGE_H,
    PAGE_W,
    _CENTER_TAB_TWIPS,
    _CONTENT_TWIPS,
    set_run_font,
    set_style_font,
)
from .footnotes import Footnotes
from .inline import parse_inline, runs_to_plain
from .math_omml import latex_math_to_text, latex_to_omml_element
from .models import (
    Equation,
    Figure,
    Heading,
    ListBlock,
    Paper,
    Para,
    Run,
    TableBlock,
)
from .style_template import load_style_document
from .util import rasterize_pdf

_RASTER_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff", ".tif", ".webp"}


class DocxBuilder:
    def __init__(
        self,
        lang: str = "ru",
        style_template: Optional[Union[str, Path]] = None,
        template: str = "note",
    ):
        self.lang = normalize_lang(lang)
        self.labels = labels_for(self.lang)
        self.profile = profile_for(template)
        self.own_typography = style_template is None
        self.footnotes = None
        if style_template:
            self.doc = load_style_document(Path(style_template))
        else:
            self.doc = docx.Document()
            section = self.doc.sections[0]
            section.page_width = PAGE_W
            section.page_height = PAGE_H
            margin = Inches(self.profile.margin_in)
            section.top_margin = section.bottom_margin = margin
            section.left_margin = section.right_margin = margin
            section.footer_distance = Inches(0.5)
            normal = self.doc.styles["Normal"]
            set_style_font(normal, self.profile.font)
            normal.font.size = Pt(self.profile.body_pt)
            normal.font.color.rgb = RGBColor(0, 0, 0)
            normal.paragraph_format.line_spacing = self.profile.line_spacing
            normal.paragraph_format.space_after = Pt(0)
            normal.paragraph_format.space_before = Pt(0)
            self._tune_heading_styles()
            self._add_page_numbers(section)

    def _font(self, run, font: Optional[str], size_pt: Optional[float]) -> None:
        """Apply Times only when this builder owns the typography.

        A style template keeps the organization's fonts. An explicit face
        such as Consolas for code is still applied.
        """
        face = self.profile.font if not font or font == FONT else font
        if not self.own_typography:
            if font and font not in (self.profile.font, FONT):
                set_run_font(run, font, size_pt=size_pt)
            return
        set_run_font(run, face, size_pt=size_pt)

    def _tune_heading_styles(self) -> None:
        for level in range(1, 4):
            style = self._find_heading_style(level)
            if style is None:
                continue
            set_style_font(style, self.profile.font)
            style.font.size = Pt(self.profile.body_pt)
            style.font.bold = True
            style.font.italic = level >= 3
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.paragraph_format.space_before = Pt(12)
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.line_spacing = 1.5
            self._set_outline(style, level - 1)

    def _find_heading_style(self, level: int):
        level = max(1, min(int(level or 1), 9))
        want_ids = {f"Heading{level}", f"heading{level}"}
        want_names = {
            f"Heading {level}",
            f"heading {level}",
            f"Заголовок {level}",
            f"заголовок {level}",
        }
        for style in self.doc.styles:
            if style.type != WD_STYLE_TYPE.PARAGRAPH:
                continue
            style_id = style.element.get(qn("w:styleId")) or ""
            if style_id in want_ids or style.name in want_names:
                return style
        return None

    @staticmethod
    def _set_outline(style, level: int) -> None:
        p_pr = style.element.find(qn("w:pPr"))
        if p_pr is None:
            p_pr = OxmlElement("w:pPr")
            style.element.append(p_pr)
        outline = p_pr.find(qn("w:outlineLvl"))
        if outline is None:
            outline = OxmlElement("w:outlineLvl")
            p_pr.append(outline)
        outline.set(qn("w:val"), str(level))

    def _add_page_numbers(self, section):
        """Centered page number in the footer."""
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run()
        set_run_font(run, self.profile.font, size_pt=max(10, self.profile.body_pt - 2))

        def _fld(char_type: str):
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), char_type)
            return el

        fld_begin = _fld("begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = " PAGE "
        fld_sep = _fld("separate")
        # visible placeholder until Word updates fields
        stub = OxmlElement("w:t")
        stub.text = "1"
        fld_end = _fld("end")

        run._r.append(fld_begin)
        run._r.append(instr)
        run._r.append(fld_sep)
        run._r.append(stub)
        run._r.append(fld_end)

    def add_p(
        self,
        text="",
        align=WD_ALIGN_PARAGRAPH.LEFT,
        first_line=0.49,
        space_before=0,
        space_after=0,
        line_spacing=1.5,
        bold=False,
        italic=False,
        size=None,
        left_indent=0,
        hanging=0,
    ):
        if size is None:
            size = self.profile.body_pt
        p = self.doc.add_paragraph()
        p.alignment = align
        if hanging > 0:
            p.paragraph_format.left_indent = Inches(left_indent + hanging)
            p.paragraph_format.first_line_indent = Inches(-hanging)
        else:
            p.paragraph_format.first_line_indent = Inches(first_line) if first_line > 0 else Inches(0)
            if left_indent:
                p.paragraph_format.left_indent = Inches(left_indent)
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        if text:
            parts = text.split("\n")
            for idx, part in enumerate(parts):
                run = p.add_run(part)
                run.bold, run.italic = bold, italic
                self._font(run, FONT, size)
                if idx < len(parts) - 1:
                    run.add_break()
        return p

    def _append_runs_to_paragraph(self, p, runs: List[Run], default_size=None):
        if default_size is None:
            default_size = BODY_PT
        for r in runs:
            if r.footnote:
                self._add_footnote(p, r.footnote)
                continue
            if r.math:
                el = latex_to_omml_element(r.math, display=False)
                if el is not None:
                    p._p.append(el)
                    continue
            parts = (r.text or "").split("\n")
            for idx, part in enumerate(parts):
                run = p.add_run(part)
                run.bold, run.italic = r.bold, r.italic
                self._font(run, r.font or FONT, r.size or default_size)
                if idx < len(parts) - 1:
                    run.add_break()

    def add_runs_para(self, para: Para):
        align_map = {
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
            "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        }
        p = self.doc.add_paragraph()
        p.alignment = align_map.get(para.align, WD_ALIGN_PARAGRAPH.JUSTIFY)
        if para.hanging > 0:
            p.paragraph_format.left_indent = Inches(para.left_indent + para.hanging)
            p.paragraph_format.first_line_indent = Inches(-para.hanging)
        else:
            p.paragraph_format.first_line_indent = Inches(para.first_line) if para.first_line > 0 else Inches(0)
            if para.left_indent:
                p.paragraph_format.left_indent = Inches(para.left_indent)
        p.paragraph_format.space_before = Pt(para.space_before)
        p.paragraph_format.space_after = Pt(para.space_after)
        p.paragraph_format.line_spacing = para.line_spacing
        default_size = para.size or BODY_PT
        self._append_runs_to_paragraph(p, para.runs, default_size=default_size)
        return p

    def add_equation(self, eq: Equation):
        """Numbered tall equations use a 2-col table (display OMML + number).

        Simple numbered eqs keep tab-stop layout with inline OMML.
        Unnumbered → centered display OMML.
        """
        tall = bool(
            re.search(
                r"\\begin\{(cases|matrix|pmatrix|bmatrix|vmatrix|Vmatrix|array|aligned|gathered)\}",
                eq.latex,
            )
            or eq.latex.count(r"\frac") >= 2
            or r"\\" in eq.latex
        )

        if eq.number is not None and tall:
            return self._add_equation_table(eq)

        p = self.doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.5

        if eq.number is not None:
            # Compact numbered: center tab + right tab
            el = latex_to_omml_element(eq.latex, display=False)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pPr = p._p.get_or_add_pPr()
            tabs = OxmlElement("w:tabs")
            tab_c = OxmlElement("w:tab")
            tab_c.set(qn("w:val"), "center")
            tab_c.set(qn("w:pos"), str(_CENTER_TAB_TWIPS))
            tab_r = OxmlElement("w:tab")
            tab_r.set(qn("w:val"), "right")
            tab_r.set(qn("w:pos"), str(_CONTENT_TWIPS))
            tabs.append(tab_c)
            tabs.append(tab_r)
            pPr.append(tabs)

            p.add_run("\t")
            if el is not None:
                p._p.append(el)
            else:
                run = p.add_run(eq.text or latex_math_to_text(eq.latex))
                run.italic = True
                self._font(run, FONT, BODY_PT)
            num_run = p.add_run(f"\t({eq.number})")
            self._font(num_run, FONT, BODY_PT)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            el_disp = latex_to_omml_element(eq.latex, display=True)
            if el_disp is not None:
                self._center_omath_para(el_disp)
                p._p.append(el_disp)
            else:
                el = latex_to_omml_element(eq.latex, display=False)
                if el is not None:
                    p._p.append(el)
                else:
                    run = p.add_run(eq.text or latex_math_to_text(eq.latex))
                    run.italic = True
                    self._font(run, FONT, BODY_PT)
        return p

    def _add_equation_table(self, eq: Equation):
        """Display OMML in left cell, (N) right-aligned — preserves tall formula size."""
        table = self.doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        cell_math, cell_num = table.rows[0].cells
        cell_math.width = Inches(5.8)
        cell_num.width = Inches(0.9)

        p = cell_math.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        el_disp = latex_to_omml_element(eq.latex, display=True)
        if el_disp is not None:
            self._center_omath_para(el_disp)
            p._p.append(el_disp)
        else:
            el = latex_to_omml_element(eq.latex, display=False)
            if el is not None:
                p._p.append(el)
            else:
                run = p.add_run(eq.text or latex_math_to_text(eq.latex))
                run.italic = True
                self._font(run, FONT, BODY_PT)

        pn = cell_num.paragraphs[0]
        pn.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pn.paragraph_format.first_line_indent = Inches(0)
        num_run = pn.add_run(f"({eq.number})")
        self._font(num_run, FONT, BODY_PT)

        # Borderless table
        for cell in (cell_math, cell_num):
            tcPr = cell._tc.get_or_add_tcPr()
            tcPr.append(
                parse_xml(
                    r'<w:tcBorders %s>'
                    r'<w:top w:val="nil"/>'
                    r'<w:left w:val="nil"/>'
                    r'<w:bottom w:val="nil"/>'
                    r'<w:right w:val="nil"/>'
                    r"</w:tcBorders>" % nsdecls("w")
                )
            )
        return table

    @staticmethod
    def _center_omath_para(el) -> None:
        """Set m:jc=center on oMathPara if present."""
        tag = el.tag if hasattr(el, "tag") else ""
        if not str(tag).endswith("oMathPara"):
            return
        pr = el.find(f"{{{OMML_NS}}}oMathParaPr")
        if pr is None:
            pr = OxmlElement("m:oMathParaPr")
            el.insert(0, pr)
        jc = pr.find(f"{{{OMML_NS}}}jc")
        if jc is None:
            jc = OxmlElement("m:jc")
            pr.append(jc)
        jc.set(qn("m:val"), "center")

    def add_heading(self, text: str, space_before=12, space_after=6, level=1):
        style = self._find_heading_style(level)
        if not self.own_typography and style is not None:
            return self.doc.add_paragraph(text, style=style)
        p = self.doc.add_paragraph()
        if style is not None:
            p.style = style
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = self.profile.line_spacing
        run = p.add_run(text)
        run.bold = True
        run.italic = level >= 3
        self._font(run, None, self.profile.body_pt)
        return p

    def _add_footnote(self, paragraph, text: str) -> None:
        if self.footnotes is None:
            self.footnotes = Footnotes(self.doc)
        self.footnotes.add(paragraph, text)

    def add_table(self, rows: List[List[List[Run]]], caption="", number=None):
        if not rows:
            return
        # TD-J: caption above the table, right-aligned
        cap = self._numbered_caption(caption, number, self.labels.table)
        if cap:
            self.add_p(
                cap,
                align=WD_ALIGN_PARAGRAPH.RIGHT,
                first_line=0,
                space_before=12,
                space_after=6,
                bold=True,
                size=11,
            )

        ncols = max(len(r) for r in rows)
        # Normalize width
        norm: List[List[List[Run]]] = []
        for r in rows:
            padded = list(r) + [[Run("")]] * (ncols - len(r))
            norm.append(padded)
        table = self.doc.add_table(rows=len(norm), cols=ncols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        col_w = Inches(6.7 / ncols)

        for row_idx, row in enumerate(table.rows):
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(r'<w:cantSplit %s/>' % nsdecls("w")))
            if row_idx == 0:
                trPr.append(parse_xml(r'<w:tblHeader %s/>' % nsdecls("w")))
            for col_idx, cell in enumerate(row.cells):
                cell.width = col_w
                # Clear default empty run, then append OMML-capable runs
                p = cell.paragraphs[0]
                p.clear()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.first_line_indent = Inches(0)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                cell_runs = norm[row_idx][col_idx]
                # Header row bold
                if row_idx == 0:
                    cell_runs = [
                        Run(
                            text=r.text,
                            bold=True,
                            italic=r.italic,
                            size=r.size or 10,
                            font=r.font,
                            math=r.math,
                            ref_key=r.ref_key,
                            eqref=r.eqref,
                        )
                        for r in cell_runs
                    ]
                else:
                    cell_runs = [
                        Run(
                            text=r.text,
                            bold=r.bold,
                            italic=r.italic,
                            size=r.size or 10,
                            font=r.font,
                            math=r.math,
                            ref_key=r.ref_key,
                            eqref=r.eqref,
                        )
                        for r in cell_runs
                    ]
                self._append_runs_to_paragraph(p, cell_runs, default_size=10)
                tcPr = cell._tc.get_or_add_tcPr()
                tcPr.append(parse_xml(
                    r'<w:tcBorders %s>'
                    r'<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
                    r'<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
                    r'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
                    r'<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
                    r"</w:tcBorders>" % nsdecls("w")
                ))

    def add_figure(self, fig: Figure):
        for path in fig.paths:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run()
            insert_path = path
            if path.exists() and path.suffix.lower() == ".pdf":
                raster = rasterize_pdf(path)
                if raster is not None:
                    insert_path = raster
            if insert_path.exists() and insert_path.suffix.lower() in _RASTER_EXTS:
                try:
                    run.add_picture(str(insert_path), width=FIG_WIDTH)
                except Exception as e:
                    run.add_text("[" + self.labels.image_failed.format(name=path.name, error=e) + "]")
            else:
                run.add_text("[" + self.labels.image_missing.format(name=path) + "]")
                run.italic = True
                self._font(run, FONT, 10)

        cap = self._numbered_caption(fig.caption, fig.number, self.labels.figure)
        if cap:
            self.add_p(cap, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=0, space_after=12, italic=True, size=10)

    def _numbered_caption(self, caption: str, number, word: str) -> str:
        cap = caption or ""
        if number is None:
            return cap
        if cap and cap.lower().startswith(CAPTION_SKIP):
            return cap
        if cap:
            return f"{word} {number}. {cap}"
        return f"{word} {number}"

    def _field(self, paper: Paper, name: str, slot: str, fallback: bool) -> str:
        value = getattr(paper, f"{name}_{slot}") or ""
        if value or not fallback:
            return value
        other = "en" if slot == "ru" else "ru"
        return getattr(paper, f"{name}_{other}") or ""

    def _add_front(self, paper: Paper, slot: str, cite_map: dict, fallback: bool) -> None:
        text_labels = labels_for(slot)
        title = self._field(paper, "title", slot, fallback)
        authors = self._field(paper, "authors", slot, fallback)
        affiliations = self._field(paper, "affiliations", slot, fallback)
        abstract = self._field(paper, "abstract", slot, fallback)
        abstract_tex = getattr(paper, f"abstract_{slot}_tex") or ""
        if fallback and not abstract_tex:
            other = "en" if slot == "ru" else "ru"
            abstract_tex = getattr(paper, f"abstract_{other}_tex") or ""
        keywords = self._field(paper, "keywords", slot, fallback)

        title_align = (
            WD_ALIGN_PARAGRAPH.CENTER if self.profile.title_align == "center" else WD_ALIGN_PARAGRAPH.LEFT
        )
        if title:
            self.add_p(title, align=title_align, first_line=0, space_after=12, bold=True, size=self.profile.body_pt)
        if authors:
            self.add_p(authors, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=0, space_after=4, bold=True)
        if affiliations:
            self.add_p(
                affiliations,
                align=WD_ALIGN_PARAGRAPH.CENTER,
                first_line=0,
                space_after=12,
                italic=True,
                size=12,
                line_spacing=1.15,
            )
        src = abstract_tex or abstract
        if src:
            if self.profile.labeled_front and not src.lower().lstrip().startswith(text_labels.abstract_skip):
                runs = [Run(text_labels.abstract + " ")] + parse_inline(src, cite_map)
            else:
                runs = parse_inline(src, cite_map)
            self.add_runs_para(self._shape_body(Para(runs=runs, align=self.profile.body_align, space_after=6)))
        if keywords:
            text = keywords
            if self.profile.labeled_front and not text.lower().startswith(text_labels.keywords_skip):
                text = text_labels.keywords + " " + text
            self.add_p(
                text,
                align=WD_ALIGN_PARAGRAPH.JUSTIFY if self.profile.body_align == "justify" else WD_ALIGN_PARAGRAPH.LEFT,
                first_line=self.profile.first_indent_in,
                space_after=18,
                size=self.profile.body_pt,
            )

    def _shape_body(self, para: Para) -> Para:
        """Apply this template's alignment and indent when the builder owns the page."""
        if not self.own_typography or para.left_indent or para.hanging:
            return para
        return Para(
            runs=para.runs,
            align=self.profile.body_align,
            first_line=self.profile.first_indent_in,
            space_before=para.space_before,
            space_after=para.space_after,
            line_spacing=self.profile.line_spacing,
            size=para.size or self.profile.body_pt,
        )

    def build(self, paper: Paper, template: str | None = None):
        if template:
            self.profile = profile_for(template)
        profile = self.profile
        cite_map = paper.cite_map or {}

        if paper.udc and profile.show_udc:
            self.add_p(f"{self.labels.udc}: {paper.udc}", first_line=0, space_after=12, bold=True)

        if profile.one_block:
            self._add_front(paper, self.lang, cite_map, fallback=True)
        else:
            self._add_front(paper, "ru", cite_map, fallback=False)
            self._add_front(paper, "en", cite_map, fallback=False)

        for block in paper.body:
            if isinstance(block, Heading):
                self.add_heading(block.text, space_before=block.space_before, space_after=block.space_after, level=block.level)
            elif isinstance(block, Para):
                self.add_runs_para(self._shape_body(block))
            elif isinstance(block, Equation):
                self.add_equation(block)
            elif isinstance(block, Figure):
                self.add_figure(block)
            elif isinstance(block, TableBlock):
                self.add_table(block.rows, caption=block.caption, number=block.number)
            elif isinstance(block, ListBlock):
                for idx, item in enumerate(block.items, 1):
                    prefix = f"{idx}. " if block.ordered else "• "
                    nest = 0.3 * block.level
                    self.add_runs_para(
                        Para(
                            runs=[Run(prefix)] + list(item),
                            align="left",
                            first_line=0 if nest else self.profile.first_indent_in,
                            left_indent=nest,
                            space_after=4,
                            line_spacing=self.profile.line_spacing,
                            size=self.profile.body_pt,
                        )
                    )

        if paper.references:
            self.add_heading(self.labels.references, space_before=18, space_after=12)
            for ref in paper.references:
                self.add_p(ref, first_line=0, space_after=4, size=12, hanging=0.3)

        if paper.authors_info:
            self.add_heading(self.labels.authors, space_before=18, space_after=12)
            for info in paper.authors_info:
                self.add_p(info, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=0.49, space_after=6, size=12)

        return self.doc

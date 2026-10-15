"""Markdown → PPTX. An organization template supplies the theme and layouts."""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional, Sequence, Union

from .util import ensure_out_prefix

PathLike = Union[str, Path]

_HR = re.compile(r"(?m)^---\s*$")
_NOTE = re.compile(r"(?m)^[?]{3}\s*$")
_HEAD = re.compile(r"^(#+)\s+(.*?)\s*#*\s*$")
_BULLET = re.compile(r"^(?:[-*+]|\d+[.)])\s+(.*)$")
_IMAGE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")


class PptxExportError(RuntimeError):
    """PPTX export failed."""


@dataclass
class SlideSpec:
    title: str = ""
    subtitle: str = ""
    lines: List[str] = field(default_factory=list)
    images: List[Path] = field(default_factory=list)
    notes: str = ""
    kind: str = "content"  # title | content


def _plain(text: str) -> str:
    s = text.strip()
    s = re.sub(r"!\[([^\]]*)\]\([^)]+\)", r"\1", s)
    s = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", s)
    s = re.sub(r"`([^`]*)`", r"\1", s)
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    s = re.sub(r"__(.+?)__", r"\1", s)
    s = re.sub(r"(?<![\w*])\*(.+?)\*(?![\w*])", r"\1", s)
    return s.strip()


def parse_slide_markdown(text: str, base: Path) -> List[SlideSpec]:
    """Slides are split on a `---` line, or on each `##` heading."""
    raw = text.replace("\r\n", "\n").strip()
    if not raw:
        return []
    if _HR.search(raw):
        parts = [p.strip() for p in _HR.split(raw) if p.strip()]
    elif re.search(r"(?m)^##\s+", raw):
        parts = [p.strip() for p in re.split(r"(?m)(?=^##\s+)", raw) if p.strip()]
    else:
        parts = [raw]
    slides = [_parse_part(part, base) for part in parts]
    return [s for s in slides if s.title or s.subtitle or s.lines or s.images or s.notes]


def _parse_part(part: str, base: Path) -> SlideSpec:
    notes = ""
    if _NOTE.search(part):
        body, notes = _NOTE.split(part, maxsplit=1)
        part = body.strip()
        notes = notes.strip()
    lines = part.splitlines()
    title = ""
    kind = "content"
    rest = lines
    if lines:
        head = _HEAD.match(lines[0].strip())
        if head:
            title = _plain(head.group(2))
            kind = "title" if len(head.group(1)) == 1 else "content"
            rest = lines[1:]
    slide = SlideSpec(title=title, kind=kind, notes=_plain(notes.replace("\n", " ")))
    body: List[str] = []
    bullets: List[str] = []
    for line in rest:
        stripped = line.strip()
        if not stripped:
            continue
        image = _IMAGE.match(stripped)
        if image:
            path = Path(image.group(2).strip())
            if not path.is_absolute():
                path = (base / path).resolve()
            slide.images.append(path)
            continue
        bullet = _BULLET.match(stripped)
        if bullet:
            bullets.append(_plain(bullet.group(1)))
            continue
        body.append(_plain(stripped))
    if kind == "title":
        slide.subtitle = "\n".join([*body, *bullets])
    else:
        if not slide.title and body:
            slide.title = body.pop(0)
        slide.lines = [*body, *bullets]
    return slide


def markdown_to_pptx(
    md_path: PathLike,
    out_path: Optional[PathLike] = None,
    *,
    template: Optional[PathLike] = None,
    title_layout: Optional[str] = None,
    content_layout: Optional[str] = None,
    lang: str = "ru",
) -> Path:
    """Build a PPTX from Markdown.

    template is an organization .pptx. Its size, theme, and layouts are kept.
    Slides already stored in that file are not copied.
    title_layout and content_layout are exact layout names from the file.
    """
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as exc:
        raise PptxExportError(
            "PPTX needs python-pptx: pip install 'latex-docx[pptx]'"
        ) from exc

    src = Path(md_path).resolve()
    if not src.exists():
        raise FileNotFoundError(src)
    slides = parse_slide_markdown(src.read_text(encoding="utf-8"), src.parent)
    if not slides:
        raise PptxExportError(f"No slides in {src.name}")

    if template:
        template_path = Path(template).resolve()
        if not template_path.exists():
            raise FileNotFoundError(template_path)
        prs = Presentation(str(template_path))
        _drop_existing_slides(prs)
    else:
        prs = Presentation()

    for spec in slides:
        layout_name = title_layout if spec.kind == "title" else content_layout
        layout = _pick_layout(prs, layout_name, spec.kind)
        slide = prs.slides.add_slide(layout)
        _fill_slide(slide, spec, Inches, lang=lang)

    dest = ensure_out_prefix(out_path) if out_path else Path(f"out_{src.stem}.pptx")
    dest.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(dest))
    return dest


def _drop_existing_slides(prs) -> None:
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst):
        prs.part.drop_rel(sld_id.rId)
        sld_id_lst.remove(sld_id)


def _placeholder_types(layout) -> List:
    return [ph.placeholder_format.type for ph in layout.placeholders]


def _pick_layout(prs, name: Optional[str], kind: str):
    from pptx.enum.shapes import PP_PLACEHOLDER as PH

    if name:
        for layout in prs.slide_layouts:
            if layout.name == name:
                return layout
        known = ", ".join(layout.name for layout in prs.slide_layouts) or "(none)"
        raise PptxExportError(f"Layout {name!r} was not found. The template has: {known}")

    title_types = {PH.TITLE, PH.CENTER_TITLE}
    body_types = {PH.BODY, PH.OBJECT}
    best = None
    best_score = -10**6
    for layout in prs.slide_layouts:
        types = _placeholder_types(layout)
        type_set = set(types)
        if not (type_set & title_types):
            continue
        bodies = sum(1 for item in types if item in body_types)
        low = layout.name.lower()
        if kind == "title":
            if bodies and PH.CENTER_TITLE not in type_set and PH.SUBTITLE not in type_set:
                continue
            score = 0
            if PH.CENTER_TITLE in type_set:
                score += 3
            if PH.SUBTITLE in type_set:
                score += 2
            if "title slide" in low or "титул" in low:
                score += 5
            if "vertical" in low or "вертик" in low:
                score -= 4
        else:
            if bodies < 1:
                continue
            score = 3 if bodies == 1 else 0
            if any(token in low for token in ("content", "содерж", "объект")):
                score += 4
            if any(token in low for token in ("vertical", "вертик", "two content", "comparison")):
                score -= 4
        if score > best_score:
            best_score = score
            best = layout
    if best is None:
        if prs.slide_layouts:
            return prs.slide_layouts[0]
        raise PptxExportError("The template has no slide layouts")
    return best


def _shape_of_type(slide, types: Sequence) -> Optional[object]:
    wanted = set(types)
    for shape in slide.placeholders:
        if shape.placeholder_format.type in wanted:
            return shape
    return None


def _write_lines(shape, lines: Sequence[str]) -> None:
    if not lines:
        return
    frame = shape.text_frame
    frame.paragraphs[0].text = lines[0]
    for extra in list(frame.paragraphs)[1:]:
        extra._p.getparent().remove(extra._p)
    for line in lines[1:]:
        frame.add_paragraph().text = line


def _fill_slide(slide, spec: SlideSpec, inches, lang: str = "ru") -> None:
    from .i18n import labels_for
    from pptx.enum.shapes import PP_PLACEHOLDER as PH

    title = _shape_of_type(slide, (PH.TITLE, PH.CENTER_TITLE))
    if title is not None and spec.title:
        _write_lines(title, [spec.title])
    subtitle = _shape_of_type(slide, (PH.SUBTITLE,))
    if subtitle is not None and spec.subtitle:
        _write_lines(subtitle, spec.subtitle.split("\n"))

    body_lines = list(spec.lines)
    if spec.subtitle and subtitle is None:
        body_lines = [*spec.subtitle.split("\n"), *body_lines]
    missing = [path.name for path in spec.images if not path.exists()]
    for name in missing:
        body_lines.append("[" + labels_for(lang).image_missing.format(name=name) + "]")
    body = _shape_of_type(slide, (PH.BODY, PH.OBJECT))
    if body is not None and body_lines:
        _write_lines(body, body_lines)
    elif body_lines and title is None:
        box = slide.shapes.add_textbox(inches(0.7), inches(0.5), inches(8), inches(1))
        _write_lines(box, [spec.title, *body_lines] if spec.title else body_lines)

    pictures = [path for path in spec.images if path.exists()]
    picture_ph = _shape_of_type(slide, (PH.PICTURE,))
    if picture_ph is not None and pictures:
        picture_ph.insert_picture(str(pictures[0]))
        pictures = pictures[1:]
    for index, path in enumerate(pictures):
        slide.shapes.add_picture(
            str(path),
            inches(0.8),
            inches(1.8 + index * 2.4),
            width=inches(6.5),
        )
    if spec.notes:
        slide.notes_slide.notes_text_frame.text = spec.notes
